#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Web API for Craveva AI Enterprise Business
Provides REST API endpoints for the React frontend
"""

from flask import Flask, request, jsonify, send_file, Response, stream_with_context
from flask_cors import CORS
import os
import sys
from datetime import datetime, date
import traceback
import re
import json
import tempfile
from io import BytesIO
import uuid
import pandas as pd
import numpy as np

def safe_json_dumps(obj, **kwargs):
    """Custom JSON serializer that handles date objects and other non-serializable types"""
    def default_serializer(o):
        if isinstance(o, (pd.Timestamp, datetime, date)):
            return o.isoformat() if hasattr(o, 'isoformat') else str(o)
        elif isinstance(o, (np.integer, np.floating)):
            return o.item()
        elif isinstance(o, np.ndarray):
            return o.tolist()
        elif pd.isna(o):
            return None
        return str(o)
    
    return json.dumps(obj, default=default_serializer, **kwargs)

# Import the analyzer
from main_mysql import E3MySQLAnalyzer
from multi_sql_flow import generate_sql_multi_decompose, ask_and_analyze_v2
# Import Excel templates
from excel_templates import enhance_prompt_with_excel_template
# Import the template system
from data_reports_templates import enhance_prompt_with_template, get_template_by_sub_option
# Import the chart generator
from chart_generator import ChartGenerator
# Import the Excel generator
from excel_generator import create_excel_file
# Import the Graphics generator
try:
    from graphics_generator import GraphicsGenerator
    GRAPHICS_AVAILABLE = True
except ImportError:
    GRAPHICS_AVAILABLE = False
    print("Warning: graphics_generator not available. Graphics export will not be available.")

# Import for DOC export
try:
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOC_EXPORT_AVAILABLE = True
except ImportError:
    DOC_EXPORT_AVAILABLE = False
    print("Warning: python-docx not installed. DOC export will not be available.")

# Import for PPT export
try:
    from ppt_generator import PPTGenerator
    PPT_EXPORT_AVAILABLE = True
except ImportError:
    PPT_EXPORT_AVAILABLE = False
    print("Warning: ppt_generator not available. PPT export will not be available.")

app = Flask(__name__)

# Configure CORS so both bare and www domains (and any others provided) can call the API
default_origins = "https://craveva.com,https://www.craveva.com"
allowed_origins = [
    origin.strip()
    for origin in os.getenv("CORS_ALLOWED_ORIGINS", default_origins).split(",")
    if origin.strip()
]
CORS(
    app,
    resources={r"/api/*": {"origins": allowed_origins}},
    supports_credentials=True
)  # Enable CORS for React frontend

# Helper function to add cache-control headers to prevent browser caching
def add_no_cache_headers(response):
    """Add cache-control headers to prevent browser caching"""
    response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '0'
    response.headers['X-Timestamp'] = str(datetime.now().timestamp())
    return response

# Initialize the analyzer
analyzer = E3MySQLAnalyzer()
# Initialize the chart generator
chart_generator = ChartGenerator()
# Initialize the PPT generator
if PPT_EXPORT_AVAILABLE:
    ppt_generator = PPTGenerator()
else:
    ppt_generator = None
# Initialize the Graphics generator
if GRAPHICS_AVAILABLE:
    graphics_generator = GraphicsGenerator()
else:
    graphics_generator = None

def auto_generate_doc_for_data_reports(session_id, question, analysis, template_id=None):
    """
    Automatically generate DOC file for Data Reports mode
    
    Args:
        session_id: Session ID
        question: User question
        analysis: AI analysis response
        template_id: Template ID for Data Reports
        
    Returns:
        dict: Contains doc_file_url and filename if successful, error message if failed
    """
    try:
        if not DOC_EXPORT_AVAILABLE:
            return {'error': 'DOC export not available. Please install python-docx.'}
        
        # Create DOC document
        doc = Document()
        
        # Add title based on template
        template_config = get_template_by_sub_option(template_id) if template_id else None
        if template_config:
            title_text = template_config['name']
        else:
            title_text = 'Data Reports Analysis'
        
        title = doc.add_heading(title_text, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add metadata
        doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"Session ID: {session_id}")
        doc.add_paragraph(f"Template: {template_config['name'] if template_config else 'General Data Reports'}")
        doc.add_paragraph("")
        
        # Add question
        doc.add_heading('Question:', level=2)
        doc.add_paragraph(question)
        doc.add_paragraph("")
        
        # Add analysis
        doc.add_heading('Analysis:', level=2)
        
        # Process analysis text, handle markdown-like formatting
        analysis_lines = analysis.split('\n')
        for line in analysis_lines:
            line = line.strip()
            if line:
                if line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                elif line.startswith('- ') or line.startswith('* '):
                    doc.add_paragraph(line[2:], style='List Bullet')
                elif line.startswith('**') and line.endswith('**'):
                    p = doc.add_paragraph()
                    run = p.add_run(line[2:-2])
                    run.bold = True
                else:
                    doc.add_paragraph(line)
        
        # Save document
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        clean_question = ''.join(c for c in question[:30] if c.isalnum() or c in (' ', '-', '_')).strip()
        filename = f"data_reports_{clean_question}_{session_id[:8]}_{timestamp}.docx"
        
        # Create exports directory if it doesn't exist
        exports_dir = 'exports'
        if not os.path.exists(exports_dir):
            os.makedirs(exports_dir)
        
        file_path = os.path.join(exports_dir, filename)
        doc.save(file_path)
        
        return {
            'doc_file_url': f'/api/download/{filename}',
            'filename': filename,
            'file_path': file_path
        }
        
    except Exception as e:
        print(f"Error generating DOC file: {e}")
        traceback.print_exc()
        return {'error': f'Failed to generate DOC file: {str(e)}'}

@app.route('/api/conversations', methods=['GET'])
def get_conversations():
    """Get list of all conversations"""
    try:
        conversations = analyzer.list_conversations()
        response = jsonify(conversations)
        return add_no_cache_headers(response)
    except Exception as e:
        print(f"Error getting conversations: {e}")
        traceback.print_exc()
        error_response = jsonify({'error': str(e), 'timestamp': datetime.now().isoformat()})
        return add_no_cache_headers(error_response), 500

@app.route('/api/conversations/<session_id>', methods=['GET'])
def get_conversation_details(session_id):
    """Get details of a specific conversation"""
    try:
        details = analyzer.get_conversation_details(session_id)
        if details:
            # Convert the ConversationContext object to a dictionary
            # Map ConversationContext attributes to frontend expected format
            
            # Extract columns from query result
            original_columns = None
            try:
                import json
                if details.query_result:
                    query_data = json.loads(details.query_result)
                    if query_data and isinstance(query_data, list) and len(query_data) > 0:
                        original_columns = list(query_data[0].keys())
            except:
                original_columns = None
            
            created_at_formatted = details.created_at
            try:
                if details.created_at:
                    created_at_formatted = datetime.fromisoformat(details.created_at).strftime('%Y-%m-%d %H:%M:%S')
            except ValueError:
                created_at_formatted = details.created_at

            updated_at_raw = getattr(details, 'updated_at', None)
            updated_at_formatted = updated_at_raw
            try:
                if updated_at_raw:
                    updated_at_formatted = datetime.fromisoformat(updated_at_raw).strftime('%Y-%m-%d %H:%M:%S')
            except ValueError:
                updated_at_formatted = updated_at_raw

            conversation_data = {
                'session_id': details.session_id,
                'original_question': details.original_question,
                'title': getattr(details, 'title', details.original_question),
                'original_sql': details.generated_sql,  # Map generated_sql to original_sql
                'original_data': details.query_result,  # Map query_result to original_data
                'original_columns': original_columns,  # Extract columns from query result
                'original_analysis': details.analysis_result,  # Map analysis_result to original_analysis
                'language': details.language,
                'created_at': details.created_at,
                'created_at_formatted': created_at_formatted,
                'updated_at': updated_at_raw,
                'updated_at_formatted': updated_at_formatted,
                'conversation_history': details.conversation_history,
                'selected_option': details.selected_option,  # Add selected_option field
                # Add chart data fields
                'has_chart': getattr(details, 'has_chart', False),
                'chart_spec': getattr(details, 'chart_spec', None),
                'chart_data': getattr(details, 'chart_data', None),
                'chart_type': getattr(details, 'chart_type', None),
                'chart_error': getattr(details, 'chart_error', None),
                # Add export data fields
                'doc_file_url': getattr(details, 'doc_file_url', None),
                'doc_filename': getattr(details, 'doc_filename', None),
                'excel_file_url': getattr(details, 'excel_file_url', None),
                'excel_filename': getattr(details, 'excel_filename', None),
                'auto_excel_generated': True if getattr(details, 'excel_file_url', None) and getattr(details, 'excel_filename', None) else False
            }
            response = jsonify(conversation_data)
            return add_no_cache_headers(response)
        else:
            error_response = jsonify({'error': 'Conversation not found', 'session_id': session_id})
            return add_no_cache_headers(error_response), 404
    except Exception as e:
        print(f"Error getting conversation details: {e}")
        traceback.print_exc()
        error_response = jsonify({'error': str(e), 'session_id': session_id, 'timestamp': datetime.now().isoformat()})
        return add_no_cache_headers(error_response), 500

@app.route('/api/conversations/<session_id>', methods=['PATCH'])
def update_conversation_metadata(session_id):
    """Update conversation metadata such as title."""
    try:
        data = request.get_json() or {}
        new_title = data.get('title')

        if new_title is None:
            return jsonify({'error': 'Title is required'}), 400

        new_title_str = str(new_title).strip()
        if not new_title_str:
            return jsonify({'error': 'Title cannot be empty'}), 400

        updated = analyzer.conversation_manager.session_manager.update_session_title(session_id, new_title_str)
        if not updated:
            return jsonify({'error': 'Conversation not found'}), 404

        details = analyzer.get_conversation_details(session_id)
        response_payload = {
            'session_id': session_id,
            'title': new_title_str
        }
        if details:
            updated_at = getattr(details, 'updated_at', None)
            response_payload['updated_at'] = updated_at
            if updated_at:
                try:
                    response_payload['updated_at_formatted'] = datetime.fromisoformat(updated_at).strftime('%Y-%m-%d %H:%M:%S')
                except ValueError:
                    response_payload['updated_at_formatted'] = updated_at
        return jsonify(response_payload)
    except Exception as e:
        print(f"Error updating conversation metadata: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/ask', methods=['POST'])
def ask_question():
    """Ask a new question"""
    try:
        data = request.get_json()
        question = data.get('question')
        session_id = data.get('session_id')  # Optional: use existing session memory if provided
        
        if not question:
            return jsonify({'error': 'Question is required'}), 400
        
        # Extract template information from question if present
        template_id = None
        enhanced_question = question
        selected_option = data.get('selected_option')  # Get from request data
        excel_type = data.get('excel_type')  # Get Excel type from request data
        
        # Check for different option patterns
        data_reports_pattern = r'\[DataReports - (\w+)\]'
        charts_pattern = r'\[Charts - (\w+)\]'
        excel_pattern = r'\[Excel - (\w+)\]'
        graphics_pattern = r'\[Graphics - (\w+)\]'
        general_pattern = r'\[(\w+)\]'
        
        # Check if question contains specific template markers
        data_reports_match = re.search(data_reports_pattern, question)
        charts_match = re.search(charts_pattern, question)
        excel_match = re.search(excel_pattern, question)
        graphics_match = re.search(graphics_pattern, question)
        general_match = re.search(general_pattern, question)
        
        chart_type = None
        
        if data_reports_match:
            template_id = data_reports_match.group(1)
            selected_option = 'DataReports'
            # Remove the marker from the question
            clean_question = re.sub(data_reports_pattern, '', question).strip()
            # Enhance the question with template-specific prompt
            enhanced_question = enhance_prompt_with_template(template_id, clean_question)
            print(f"Using template: {template_id}")
            print(f"Enhanced question: {enhanced_question}")
        elif charts_match:
            chart_type = charts_match.group(1)
            selected_option = 'Charts'
            # Remove the marker from the question
            enhanced_question = re.sub(charts_pattern, '', question).strip()
            print(f"Using Charts mode with {chart_type} chart type")
        elif graphics_match:
            graphics_type = graphics_match.group(1)
            selected_option = 'Graphics'
            # Remove the marker from the question
            enhanced_question = re.sub(graphics_pattern, '', question).strip()
            print(f"Using Graphics mode with {graphics_type} graphics type")
        elif excel_match:
            excel_type = excel_match.group(1)
            selected_option = 'Excel'
            # Remove the marker from the question
            clean_question = re.sub(excel_pattern, '', question).strip()
            # Enhance the question with Excel template-specific prompt
            enhanced_question = enhance_prompt_with_excel_template(excel_type, clean_question)
            print(f"Using Excel mode with {excel_type} type")
            print(f"Enhanced question: {enhanced_question}")
        elif general_match:
            selected_option = general_match.group(1)
            # Remove the marker from the question
            enhanced_question = re.sub(general_pattern, '', question).strip()
            print(f"Using {selected_option} mode")
        
        # Handle direct parameter input for Excel mode
        if selected_option == 'Excel' and excel_type and not excel_match:
            # Enhance the question with Excel template-specific prompt
            enhanced_question = enhance_prompt_with_excel_template(excel_type, question)
            print(f"Using Excel mode with {excel_type} type (direct parameter)")
            print(f"Enhanced question: {enhanced_question}")
        
        # Prepare chart data for Charts mode (before calling ask_and_analyze)
        chart_data_for_session = None
        chart_spec_for_session = None
        has_chart_for_session = False
        chart_error_for_session = None
        
        # Prepare graphics data for Graphics mode (before calling ask_and_analyze)
        graphics_file_url_for_session = None
        graphics_filename_for_session = None
        graphics_error_for_session = None
        
        # For Charts mode, we need to first get the data to generate charts
        if selected_option == 'Charts' and chart_type:
            # First, get the data without creating a session
            temp_query_data, temp_analysis, _ = analyzer.ask_and_analyze(
                enhanced_question,
                create_session=False,
                selected_option=selected_option,
                session_id=session_id  # inject memory if user provided an existing session
            )
            
            if temp_query_data is not None:
                try:
                    # Try to generate chart from query data
                    if (hasattr(temp_query_data, 'empty') and not temp_query_data.empty) or \
                       (isinstance(temp_query_data, list) and len(temp_query_data) > 0):
                        # Convert query data to list of dictionaries if needed
                        if hasattr(temp_query_data, 'to_dict'):
                            chart_data_for_session = temp_query_data.to_dict('records')
                        elif isinstance(temp_query_data, list):
                            chart_data_for_session = temp_query_data
                        else:
                            chart_data_for_session = []
                        
                        # Generate chart specification
                        chart_spec_for_session = chart_generator.generate_chart(
                            data=chart_data_for_session,
                            chart_type=chart_type
                        )
                        
                        has_chart_for_session = True
                    else:
                        chart_error_for_session = "No data available for chart generation"
                        
                except Exception as chart_error:
                    print(f"Chart generation error: {chart_error}")
                    chart_error_for_session = str(chart_error)
        
        # Pre-generate Graphics file for Graphics mode (before creating session)
        elif selected_option == 'Graphics' and graphics_type and GRAPHICS_AVAILABLE:
            # We need to generate a temporary session ID for Graphics generation
            temp_session_id = str(uuid.uuid4())[:8]
            
            # First, get a quick analysis to use for Graphics generation
            temp_query_data, temp_analysis, _ = analyzer.ask_and_analyze(
                enhanced_question, 
                create_session=False,  # Don't create session yet
                selected_option=selected_option,
                session_id=session_id  # inject memory if provided
            )
            
            if temp_analysis and temp_query_data is not None:
                try:
                    # Convert query data to list of dictionaries if needed
                    if hasattr(temp_query_data, 'to_dict'):
                        graphics_data = temp_query_data.to_dict('records')
                    elif isinstance(temp_query_data, list):
                        graphics_data = temp_query_data
                    else:
                        graphics_data = []
                    
                    # Generate Graphics file
                    graphics_filepath = graphics_generator.generate_graphics_by_type(
                        graphics_type=graphics_type,
                        data=graphics_data,
                        analysis_text=temp_analysis,
                        title=f"{graphics_type.title()} Analysis"
                    )
                    
                    if graphics_filepath and os.path.exists(graphics_filepath):
                        graphics_filename = os.path.basename(graphics_filepath)
                        graphics_file_url_for_session = f'/api/download/{graphics_filename}'
                        graphics_filename_for_session = graphics_filename
                        
                except Exception as graphics_error:
                    print(f"Graphics generation error: {graphics_error}")
                    graphics_error_for_session = str(graphics_error)
        
        # Initialize DOC-related variables
        doc_file_url_for_session = None
        doc_filename_for_session = None
        
        # Initialize Excel-related variables
        excel_file_url_for_session = None
        excel_filename_for_session = None
        
        # Pre-generate DOC file for DataReports mode (before creating session)
        if selected_option == 'DataReports':
            # We need to generate a temporary session ID for DOC generation
            temp_session_id = str(uuid.uuid4())[:8]
            
            # First, get a quick analysis to use for DOC generation
            temp_query_data, temp_analysis, _ = analyzer.ask_and_analyze(
                enhanced_question, 
                create_session=False,  # Don't create session yet
                selected_option=selected_option,
                session_id=session_id  # inject memory if provided
            )
            
            if temp_analysis:
                # Generate DOC file with temporary data
                doc_result = auto_generate_doc_for_data_reports(temp_session_id, enhanced_question, temp_analysis, template_id)
                
                if 'error' not in doc_result:
                    doc_file_url_for_session = doc_result['doc_file_url']
                    doc_filename_for_session = doc_result['filename']
        
        # Pre-generate Excel file for Excel mode (before creating session)
        elif selected_option == 'Excel' and excel_type:
            # We need to generate a temporary session ID for Excel generation
            temp_session_id = str(uuid.uuid4())[:8]
            
            # First, get a quick analysis to use for Excel generation
            temp_query_data, temp_analysis, _ = analyzer.ask_and_analyze(
                enhanced_question, 
                create_session=False,  # Don't create session yet
                selected_option=selected_option,
                session_id=session_id  # inject memory if provided
            )
            
            if temp_analysis:
                try:
                    # Generate Excel file with temporary data
                    excel_filepath, excel_filename = create_excel_file(
                        excel_type=excel_type,
                        data=temp_query_data if temp_query_data is not None else "No data available",
                        analysis_text=temp_analysis,
                        session_id=temp_session_id
                    )
                    
                    # Create download URL
                    excel_file_url_for_session = f"/api/download/{excel_filename}"
                    excel_filename_for_session = excel_filename
                    print(f"Excel file generated: {excel_filename}")
                    
                except Exception as excel_error:
                    print(f"Excel generation error: {excel_error}")
                    # Continue without Excel file if generation fails

        # Use the LLM-first multi-SQL flow to process the question (with chart, DOC, and Excel data if available)
        query_data, analysis, session_id = ask_and_analyze_v2(
            analyzer,
            enhanced_question, 
            create_session=True, 
            selected_option=selected_option,
            has_chart=has_chart_for_session,
            chart_spec=chart_spec_for_session,
            chart_data=chart_data_for_session,
            chart_type=chart_type,
            chart_error=chart_error_for_session,
            doc_file_url=doc_file_url_for_session,
            doc_filename=doc_filename_for_session,
            excel_file_url=excel_file_url_for_session,
            excel_filename=excel_filename_for_session,
            session_id=session_id
        )
        
        if query_data is None:
            return jsonify({'error': 'Failed to process question'}), 500
        
        response = {
            'session_id': session_id,
            'question': question,
            'analysis': analysis,
            'timestamp': datetime.now().isoformat(),
            'template_id': template_id,
            'template_config': get_template_by_sub_option(template_id) if template_id else None,
            'selected_option': selected_option,
            'chart_type': chart_type
        }
        # Include multi-SQL metadata if available
        try:
            sql_list = getattr(analyzer, 'last_sql_list', None)
            result_sets = getattr(analyzer, 'last_result_sets', None)
            if sql_list:
                response['sql_list'] = sql_list
            if result_sets:
                response['result_sets'] = result_sets
        except Exception:
            # Non-critical; ignore if analyzer does not provide these attributes
            pass
        
        # Add chart data to response if available
        if has_chart_for_session:
            response['chart_spec'] = chart_spec_for_session
            response['chart_data'] = chart_data_for_session
            response['has_chart'] = True
            
            # Add chart info to analysis
            chart_info_msg = f"\n\n---\n\n**📊 {chart_type} Chart Generated!**\n\nA {chart_type.lower()} chart has been generated based on your query data."
            response['analysis'] = analysis + chart_info_msg
        elif chart_error_for_session:
            response['has_chart'] = False
            response['chart_error'] = chart_error_for_session
        
        # Add DOC data to response if available
        if doc_file_url_for_session and doc_filename_for_session:
            # Add download link to analysis
            doc_success_msg = f"\n\n---\n\n**📄 DOC Report Generated Successfully!**\n\nYour comprehensive data analysis report has been automatically generated and is ready for download.\n\n**File:** {doc_filename_for_session}\n\n[Download DOC Report]({doc_file_url_for_session})"
            response['analysis'] = analysis + doc_success_msg
            response['doc_file_url'] = doc_file_url_for_session
            response['doc_filename'] = doc_filename_for_session
            response['auto_doc_generated'] = True
        
        # Add Excel data to response if available
        if excel_file_url_for_session and excel_filename_for_session:
            # Add download link to analysis
            excel_success_msg = f"\n\n---\n\n**📊 Excel Report Generated Successfully!**\n\nYour comprehensive Excel analysis report has been automatically generated and is ready for download.\n\n**File:** {excel_filename_for_session}\n\n[Download Excel Report]({excel_file_url_for_session})"
            response['analysis'] = analysis + excel_success_msg
            response['excel_file_url'] = excel_file_url_for_session
            response['excel_filename'] = excel_filename_for_session
            response['auto_excel_generated'] = True
        
        return jsonify(response)
        
    except Exception as e:
        print(f"Error processing question: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/ask/stream', methods=['GET'])
def ask_question_stream():
    """Stream-first ask endpoint using Server-Sent Events (SSE).
    Streams: start, sql_list, execution_progress, analysis_token, analysis_final, done
    """
    try:
        question = request.args.get('question', '').strip()
        # Use a distinct variable for incoming session to avoid shadowing inside generator
        incoming_session_id = request.args.get('session_id', '').strip() or None
        if not question:
            return jsonify({'error': 'Question is required'}), 400

        # Extract template markers similar to /api/ask
        template_id = None
        enhanced_question = question
        selected_option = None
        chart_type = None
        excel_type = None

        data_reports_pattern = r'\[DataReports - (\w+)\]'
        charts_pattern = r'\[Charts - (\w+)\]'
        excel_pattern = r'\[Excel - (\w+)\]'
        flash_pattern = r'\[Flash\]'
        general_pattern = r'\[(\w+)\]'

        data_reports_match = re.search(data_reports_pattern, question)
        charts_match = re.search(charts_pattern, question)
        excel_match = re.search(excel_pattern, question)
        flash_match = re.search(flash_pattern, question)
        general_match = re.search(general_pattern, question)

        if data_reports_match:
            template_id = data_reports_match.group(1)
            selected_option = 'DataReports'
            clean_question = re.sub(data_reports_pattern, '', question).strip()
            enhanced_question = enhance_prompt_with_template(template_id, clean_question)
        elif charts_match:
            chart_type = charts_match.group(1)
            selected_option = 'Charts'
            enhanced_question = re.sub(charts_pattern, '', question).strip()
        elif excel_match:
            excel_type = excel_match.group(1)
            selected_option = 'Excel'
            clean_question = re.sub(excel_pattern, '', question).strip()
            enhanced_question = enhance_prompt_with_excel_template(excel_type, clean_question)
        elif flash_match:
            selected_option = 'Flash'
            enhanced_question = re.sub(flash_pattern, '', question).strip()
        elif general_match:
            selected_option = general_match.group(1)
            enhanced_question = re.sub(general_pattern, '', question).strip()

        def sse(event, payload):
            return f"event: {event}\n" + f"data: {safe_json_dumps(payload)}\n\n"

        @stream_with_context
        def generate():
            try:
                # Announce start
                yield sse('start', {
                    'question': question,
                    'selected_option': selected_option,
                    'template_id': template_id,
                    'chart_type': chart_type,
                    'excel_type': excel_type
                })

                # Fast path for Flash mode: single SQL and lightweight analysis
                if selected_option == 'Flash':
                    # Pass through any incoming session context, but do not reassign the name inside generator
                    sql_stmt = analyzer.generate_sql(enhanced_question, mode=selected_option, session_id=incoming_session_id)
                    if not sql_stmt:
                        yield sse('analysis_final', {'text': 'Unable to generate SQL for this question.'})
                        yield sse('done', {'error': 'No SQL generated'})
                        return

                    yield sse('sql_list', {
                        'items': [{
                            'name': 'Query 1',
                            'purpose': 'Single query',
                            'sql': sql_stmt
                        }]
                    })

                    # Handle both string and dict returns from generate_sql
                    if isinstance(sql_stmt, dict) and 'sql' in sql_stmt:
                        sql_exec = (sql_stmt['sql'] or '').strip().rstrip(';')
                    else:
                        sql_exec = (sql_stmt or '').strip().rstrip(';')
                    df = analyzer.execute_sql(sql_exec)
                    rows = 0
                    cols = []
                    if df is not None:
                        try:
                            rows = int(df.shape[0])
                            cols = list(df.columns)
                        except Exception:
                            rows = 0
                            cols = []

                    yield sse('execution_progress', {
                        'index': 1,
                        'name': 'Query 1',
                        'rows': rows,
                        'columns': cols
                    })

                    datasets = [{
                        'name': 'Query 1',
                        'purpose': 'Single query',
                        'sql': sql_exec,
                        'df': df
                    }]

                    # Stream analysis tokens using lightweight Flash prompt
                    accumulated = ''
                    for token in analyzer.analyze_multi_data(datasets, enhanced_question, mode=selected_option, stream=True, session_id=incoming_session_id):
                        if token:
                            accumulated += token
                            yield sse('analysis_token', {'text': token})
                    yield sse('analysis_final', {'text': accumulated})

                    # Persist conversation
                    primary_df = df if (df is not None and getattr(df, 'empty', False) is False) else (df if df is not None else pd.DataFrame())
                    # Persist the conversation and use a distinct variable to avoid shadowing
                    # Save conversation to database with error handling
                    persisted_session_id = None
                    try:
                        persisted_session_id = analyzer.conversation_manager.create_conversation(
                            question=question,
                            sql_query=sql_exec,
                            query_result=primary_df.to_dict('records') if hasattr(primary_df, 'to_dict') else [],
                            columns=list(primary_df.columns) if hasattr(primary_df, 'columns') else [],
                            analysis=accumulated,
                            selected_option=selected_option,
                            has_chart=False,
                            chart_spec=None,
                            chart_data=None,
                            chart_type=None,
                            chart_error=None,
                            doc_file_url=None,
                            doc_filename=None,
                            excel_file_url=None,
                            excel_filename=None,
                            user_id="6"
                        )
                        print(f"✅ Conversation saved successfully: session_id={persisted_session_id}")
                    except Exception as db_error:
                        print(f"❌ CRITICAL: Failed to save conversation to database: {db_error}")
                        traceback.print_exc()
                        # DO NOT return temp session_id - return error instead
                        # This prevents data loss - frontend must handle error and retry
                        yield sse('done', {
                            'error': 'Failed to save conversation to database. Please try again.',
                            'session_id': None,
                            'selected_option': selected_option,
                            'template_id': template_id,
                            'chart_type': chart_type,
                            'has_chart': False,
                            'chart_spec': None,
                            'chart_data': None,
                            'chart_error': None,
                            'sql_list': [sql_exec],
                            'result_sets': [{
                                'name': 'Query 1',
                                'purpose': 'Single query',
                                'rows': rows,
                                'columns': cols
                            }],
                            'doc_file_url': None,
                            'doc_filename': None,
                            'excel_file_url': None,
                            'excel_filename': None
                        })
                        return  # Exit early - don't continue with unsaved conversation

                    yield sse('done', {
                        'session_id': persisted_session_id,
                        'selected_option': selected_option,
                        'template_id': template_id,
                        'chart_type': chart_type,
                        'has_chart': False,
                        'chart_spec': None,
                        'chart_data': None,
                        'chart_error': None,
                        'sql_list': [sql_exec],
                        'result_sets': [{
                            'name': 'Query 1',
                            'purpose': 'Single query',
                            'rows': rows,
                            'columns': cols
                        }],
                        'doc_file_url': None,
                        'doc_filename': None,
                        'excel_file_url': None,
                        'excel_filename': None
                    })
                    return

                # Generate SQL list via LLM-first decomposition
                sql_items = generate_sql_multi_decompose(analyzer, enhanced_question, mode=selected_option, session_id=incoming_session_id)
                if not sql_items:
                    yield sse('analysis_final', {'text': 'Unable to generate SQL for this question.'})
                    yield sse('done', {'error': 'No SQL generated'})
                    return
                yield sse('sql_list', {
                    'items': [{'name': i.get('name'), 'purpose': i.get('purpose'), 'sql': i.get('sql')} for i in sql_items]
                })

                # Execute SQLs in parallel and stream progress
                from concurrent.futures import ThreadPoolExecutor, as_completed
                datasets = []
                
                def execute_single_query(idx, item):
                    """Execute a single SQL query (thread-safe)"""
                    sql_stmt = (item.get('sql') or '').strip().rstrip(';')
                    if not sql_stmt:
                        return None, idx, item
                    print(f"Executing SQL {idx}/{len(sql_items)}: {sql_stmt[:100]}...")
                    try:
                        # Use thread connection for parallel execution
                        df = analyzer.execute_sql(sql_stmt, use_thread_connection=True)
                        return df, idx, item
                    except Exception as e:
                        print(f"❌ Error executing query {idx}: {e}")
                        return None, idx, item
                
                # Execute queries in parallel (max 5 concurrent)
                with ThreadPoolExecutor(max_workers=5) as executor:
                    future_to_idx = {
                        executor.submit(execute_single_query, idx, item): idx 
                        for idx, item in enumerate(sql_items, start=1)
                    }
                    completed_results = {}
                    for future in as_completed(future_to_idx):
                        df, idx, item = future.result()
                        completed_results[idx] = (df, item)
                
                # Process results in original order and stream progress
                for idx in sorted(completed_results.keys()):
                    df, item = completed_results[idx]
                    sql_stmt = (item.get('sql') or '').strip().rstrip(';')
                    rows = 0
                    cols = []
                    if df is not None:
                        try:
                            rows = int(df.shape[0])
                            cols = list(df.columns)
                            if rows == 0:
                                print(f"⚠️ Query {idx} returned empty result")
                            else:
                                print(f"✅ Query {idx} returned {rows} rows")
                        except Exception as e:
                            print(f"⚠️ Error processing DataFrame for query {idx}: {e}")
                            rows = 0
                            cols = []
                        # Add dataset even if empty (for analysis to report "no data")
                        datasets.append({
                            'name': item.get('name'),
                            'purpose': item.get('purpose'),
                            'sql': sql_stmt,
                            'df': df
                        })
                    else:
                        print(f"❌ Query {idx} execution failed (returned None)")
                    yield sse('execution_progress', {
                        'index': idx,
                        'name': item.get('name') or f'Dataset {idx}',
                        'rows': rows,
                        'columns': cols
                    })

                # Check if we have any datasets at all
                if not datasets:
                    print("❌ No valid datasets generated - all queries failed")
                    yield sse('analysis_final', {'text': 'No data found for the generated queries. All SQL queries either failed or returned no results.'})
                    yield sse('done', {'error': 'No datasets'})
                    return
                
                # Check if all datasets are empty
                all_empty = True
                for d in datasets:
                    if d['df'] is not None:
                        try:
                            if not d['df'].empty:
                                all_empty = False
                                break
                        except Exception:
                            pass
                
                if all_empty:
                    print("⚠️ All queries returned empty results")
                    # Still proceed with analysis so AI can report "no data found"

                # Stream analysis tokens
                accumulated = ''
                for token in analyzer.analyze_multi_data(datasets, enhanced_question, mode=selected_option, stream=True, session_id=incoming_session_id):
                    if token:
                        accumulated += token
                        yield sse('analysis_token', {'text': token})
                yield sse('analysis_final', {'text': accumulated})

                # Generate chart for Charts mode
                chart_spec_for_session = None
                chart_data_for_session = None
                has_chart_for_session = False
                chart_error_for_session = None
                
                if selected_option == 'Charts' and chart_type:
                    # Choose primary df for chart generation
                    primary_df = None
                    for d in datasets:
                        if d['df'] is not None and getattr(d['df'], 'empty', False) is False:
                            primary_df = d['df']
                            break
                    if primary_df is None and datasets:
                        primary_df = datasets[0]['df']
                    
                    if primary_df is not None:
                        try:
                            # Convert query data to list of dictionaries for chart generation
                            if hasattr(primary_df, 'to_dict'):
                                chart_data_for_session = primary_df.to_dict('records')
                            elif isinstance(primary_df, list):
                                chart_data_for_session = primary_df
                            else:
                                chart_data_for_session = []
                            
                            # Generate chart specification
                            if chart_data_for_session:
                                chart_spec_for_session = chart_generator.generate_chart(
                                    data=chart_data_for_session,
                                    chart_type=chart_type
                                )
                                has_chart_for_session = True
                            else:
                                chart_error_for_session = "No data available for chart generation"
                                
                        except Exception as chart_error:
                            print(f"Chart generation error: {chart_error}")
                            chart_error_for_session = str(chart_error)

                # Persist conversation
                # Choose primary df compatible with existing storage
                primary_df = None
                for d in datasets:
                    if d['df'] is not None and getattr(d['df'], 'empty', False) is False:
                        primary_df = d['df']
                        break
                if primary_df is None:
                    primary_df = datasets[0]['df']

                joined_sql = "\n\n".join([f"-- {i+1}. {item.get('name','Query')}\n{(item.get('sql') or '').strip()}" for i, item in enumerate(datasets)])
                # Save conversation to database with error handling
                persisted_session_id = None
                try:
                    persisted_session_id = analyzer.conversation_manager.create_conversation(
                        question=question,
                        sql_query=joined_sql,
                        query_result=primary_df.to_dict('records') if hasattr(primary_df, 'to_dict') else [],
                        columns=list(primary_df.columns) if hasattr(primary_df, 'columns') else [],
                        analysis=accumulated,
                        selected_option=selected_option,
                        has_chart=has_chart_for_session,
                        chart_spec=chart_spec_for_session,
                        chart_data=chart_data_for_session,
                        chart_type=chart_type,
                        chart_error=chart_error_for_session,
                        graphics_file_url=graphics_file_url_for_session,
                        graphics_filename=graphics_filename_for_session,
                        graphics_error=graphics_error_for_session,
                        doc_file_url=None,
                        doc_filename=None,
                        excel_file_url=None,
                        excel_filename=None,
                        user_id="6"
                    )
                    print(f"✅ Conversation saved successfully: session_id={persisted_session_id}")
                except Exception as db_error:
                    print(f"❌ CRITICAL: Failed to save conversation to database: {db_error}")
                    traceback.print_exc()
                    # DO NOT return temp session_id - return error instead
                    # This prevents data loss - frontend must handle error and retry
                    yield sse('done', {
                        'error': 'Failed to save conversation to database. Please try again.',
                        'session_id': None,
                        'selected_option': selected_option,
                        'template_id': template_id,
                        'chart_type': chart_type,
                        'has_chart': has_chart_for_session,
                        'chart_spec': chart_spec_for_session,
                        'chart_data': chart_data_for_session,
                        'chart_error': chart_error_for_session,
                        'graphics_file_url': graphics_file_url_for_session,
                        'graphics_filename': graphics_filename_for_session,
                        'graphics_error': graphics_error_for_session,
                        'sql_list': [d['sql'] for d in datasets],
                        'result_sets': [{
                            'name': d.get('name'),
                            'purpose': d.get('purpose'),
                            'rows': int(d['df'].shape[0]) if hasattr(d['df'], 'shape') else 0,
                            'columns': list(d['df'].columns) if hasattr(d['df'], 'columns') else []
                        } for d in datasets],
                        'doc_file_url': None,
                        'doc_filename': None,
                        'excel_file_url': None,
                        'excel_filename': None
                    })
                    return  # Exit early - don't continue with unsaved conversation

                # Optional: generate DOC/Excel after analysis for DataReports/Excel modes
                doc_file_url = None
                doc_filename = None
                excel_file_url = None
                excel_filename = None
                if selected_option == 'DataReports' and template_id:
                    try:
                        doc_result = auto_generate_doc_for_data_reports(persisted_session_id, enhanced_question, accumulated, template_id)
                        if 'error' not in doc_result:
                            doc_file_url = doc_result.get('doc_file_url')
                            doc_filename = doc_result.get('filename')
                            # Persist DOC info to conversation so it survives refresh
                            try:
                                context = analyzer.conversation_manager.get_conversation_details(persisted_session_id)
                                if context:
                                    context.doc_file_url = doc_file_url
                                    context.doc_filename = doc_filename
                                    analyzer.conversation_manager.session_manager.update_session(persisted_session_id, context)
                            except Exception:
                                pass
                    except Exception as _:
                        pass
                if selected_option == 'Excel' and excel_type:
                    try:
                        excel_path, excel_name = create_excel_file(
                            excel_type=excel_type,
                            data=primary_df if primary_df is not None else [],
                            analysis_text=accumulated,
                            session_id=persisted_session_id
                        )
                        excel_file_url = f"/api/download/{excel_name}"
                        excel_filename = excel_name
                        # Persist Excel info to conversation so it survives refresh
                        try:
                            context = analyzer.conversation_manager.get_conversation_details(persisted_session_id)
                            if context:
                                context.excel_file_url = excel_file_url
                                context.excel_filename = excel_filename
                                analyzer.conversation_manager.session_manager.update_session(persisted_session_id, context)
                        except Exception:
                            pass
                    except Exception as _:
                        pass

                yield sse('done', {
                    'session_id': persisted_session_id,
                    'selected_option': selected_option,
                    'template_id': template_id,
                    'chart_type': chart_type,
                    'has_chart': has_chart_for_session,
                    'chart_spec': chart_spec_for_session,
                    'chart_data': chart_data_for_session,
                    'chart_error': chart_error_for_session,
                    'sql_list': [d['sql'] for d in datasets],
                    'result_sets': [{
                        'name': d.get('name'),
                        'purpose': d.get('purpose'),
                        'rows': int(d['df'].shape[0]) if hasattr(d['df'], 'shape') else 0,
                        'columns': list(d['df'].columns) if hasattr(d['df'], 'columns') else []
                    } for d in datasets],
                    'doc_file_url': doc_file_url,
                    'doc_filename': doc_filename,
                    'excel_file_url': excel_file_url,
                    'excel_filename': excel_filename
                })
            except Exception as e:
                yield sse('done', {'error': str(e)})

        resp = Response(generate(), mimetype='text/event-stream')
        # Disable buffering for proxies like nginx and enable CORS/keep-alive for SSE
        resp.headers['Cache-Control'] = 'no-cache'
        resp.headers['X-Accel-Buffering'] = 'no'
        resp.headers['Connection'] = 'keep-alive'
        resp.headers['Access-Control-Allow-Origin'] = '*'
        resp.headers['Access-Control-Allow-Headers'] = 'Content-Type'
        return resp
    except Exception as e:
        print(f"Error in /api/ask/stream: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/followup/stream', methods=['GET'])
def followup_question_stream():
    """Stream follow-up analysis via SSE - now uses complete data analysis workflow"""
    try:
        session_id = request.args.get('session_id', '').strip()
        question = request.args.get('question', '').strip()

        if not session_id or not question:
            return jsonify({'error': 'Session ID and question are required'}), 400

        # CRITICAL: Reject temporary session IDs to prevent data loss
        if session_id.startswith('temp_') or session_id.startswith('temp-'):
            return jsonify({
                'error': 'Invalid session ID. The previous conversation was not saved. Please start a new conversation.',
                'session_id': None
            }), 400

        # Extract template markers similar to /api/ask/stream
        template_id = None
        enhanced_question = question
        selected_option = None
        chart_type = None
        excel_type = None

        data_reports_pattern = r'\[DataReports - (\w+)\]'
        charts_pattern = r'\[Charts - (\w+)\]'
        excel_pattern = r'\[Excel - (\w+)\]'
        flash_pattern = r'\[Flash\]'
        general_pattern = r'\[(\w+)\]'

        data_reports_match = re.search(data_reports_pattern, question)
        charts_match = re.search(charts_pattern, question)
        excel_match = re.search(excel_pattern, question)
        flash_match = re.search(flash_pattern, question)
        general_match = re.search(general_pattern, question)

        if data_reports_match:
            template_id = data_reports_match.group(1)
            selected_option = 'DataReports'
            clean_question = re.sub(data_reports_pattern, '', question).strip()
            enhanced_question = enhance_prompt_with_template(template_id, clean_question)
        elif charts_match:
            chart_type = charts_match.group(1)
            selected_option = 'Charts'
            enhanced_question = re.sub(charts_pattern, '', question).strip()
        elif excel_match:
            excel_type = excel_match.group(1)
            selected_option = 'Excel'
            clean_question = re.sub(excel_pattern, '', question).strip()
            enhanced_question = enhance_prompt_with_excel_template(excel_type, clean_question)
        elif flash_match:
            selected_option = 'Flash'
            enhanced_question = re.sub(flash_pattern, '', question).strip()
        elif general_match:
            selected_option = general_match.group(1)
            enhanced_question = re.sub(general_pattern, '', question).strip()

        def sse(event, payload):
            return f"event: {event}\n" + f"data: {safe_json_dumps(payload)}\n\n"

        @stream_with_context
        def generate():
            try:
                # Get language from existing conversation context
                context = analyzer.conversation_manager.get_conversation_details(session_id)
                if not context:
                    yield sse('analysis_final', {'text': 'Conversation not found.'})
                    yield sse('done', {'error': 'Conversation not found'})
                    return
                
                language = context.language

                # Announce start
                yield sse('start', {
                    'session_id': session_id,
                    'question': question,
                    'selected_option': selected_option,
                    'template_id': template_id,
                    'chart_type': chart_type,
                    'excel_type': excel_type
                })

                # Fast path for Flash followup: single SQL and lightweight analysis
                if selected_option == 'Flash':
                    sql_stmt = analyzer.generate_sql(enhanced_question, mode=selected_option, session_id=session_id)
                    if not sql_stmt:
                        yield sse('analysis_final', {'text': 'Unable to generate SQL for this question.'})
                        yield sse('done', {'error': 'No SQL generated'})
                        return

                    yield sse('sql_list', {
                        'items': [{
                            'name': 'Query 1',
                            'purpose': 'Single query',
                            'sql': sql_stmt
                        }]
                    })

                    sql_exec = (sql_stmt or '').strip().rstrip(';')
                    df = analyzer.execute_sql(sql_exec)
                    rows = 0
                    cols = []
                    if df is not None:
                        try:
                            rows = int(df.shape[0])
                            cols = list(df.columns)
                        except Exception:
                            rows = 0
                            cols = []

                    yield sse('execution_progress', {
                        'index': 1,
                        'name': 'Query 1',
                        'rows': rows,
                        'columns': cols
                    })

                    datasets = [{
                        'name': 'Query 1',
                        'purpose': 'Single query',
                        'sql': sql_exec,
                        'df': df
                    }]

                    accumulated = ''
                    for token in analyzer.analyze_multi_data(datasets, enhanced_question, mode=selected_option, stream=True, session_id=session_id):
                        if token:
                            accumulated += token
                            yield sse('analysis_token', {'text': token})
                    yield sse('analysis_final', {'text': accumulated})

                    # Add to existing conversation as followup
                    try:
                        # Get current conversation context
                        context = analyzer.conversation_manager.get_conversation_details(session_id)
                        if not context:
                            print(f"⚠️ Warning: Session {session_id} not found, cannot save followup")
                        else:
                            primary_df = df if (df is not None and getattr(df, 'empty', False) is False) else (df if df is not None else pd.DataFrame())
                            joined_sql = sql_exec
                            followup_data = {
                                'timestamp': datetime.now().isoformat(),
                                'question': question,
                                'sql_query': joined_sql,
                                'query_result': primary_df.to_dict('records') if hasattr(primary_df, 'to_dict') else [],
                                'columns': list(primary_df.columns) if hasattr(primary_df, 'columns') else [],
                                'analysis': accumulated,
                                'type': 'followup',
                                'selected_option': selected_option
                            }
                            if context.conversation_history is None:
                                context.conversation_history = []
                            context.conversation_history.append(followup_data)
                            analyzer.conversation_manager.session_manager.update_session(session_id, context)
                            print(f"✅ Successfully saved followup to conversation session_id={session_id}")
                    except Exception as e:
                        print(f"❌ Error saving followup to database: {e}")
                        print(f"   Session ID: {session_id}")
                        traceback.print_exc()
                        # Don't fail the entire request, but log the error

                    yield sse('done', {
                        'session_id': session_id,
                        'selected_option': selected_option,
                        'template_id': template_id,
                        'chart_type': chart_type,
                        'excel_type': excel_type,
                        'has_chart': False,
                        'chart_spec': None,
                        'chart_data': None,
                        'chart_error': None
                    })
                    return

                # Generate SQL list (same as initial query)
                sql_items = generate_sql_multi_decompose(analyzer, enhanced_question, mode=selected_option, session_id=session_id)
                if not sql_items:
                    yield sse('analysis_final', {'text': 'Unable to generate SQL for this question.'})
                    yield sse('done', {'error': 'No SQL generated'})
                    return
                yield sse('sql_list', {
                    'items': [{'name': i.get('name'), 'purpose': i.get('purpose'), 'sql': i.get('sql')} for i in sql_items]
                })

                # Execute SQLs and stream progress (same as initial query)
                # Execute SQLs in parallel and stream progress
                from concurrent.futures import ThreadPoolExecutor, as_completed
                datasets = []
                
                def execute_single_query(idx, item):
                    """Execute a single SQL query (thread-safe)"""
                    sql_stmt = (item.get('sql') or '').strip().rstrip(';')
                    if not sql_stmt:
                        return None, idx, item
                    print(f"Executing SQL {idx}/{len(sql_items)}: {sql_stmt[:100]}...")
                    try:
                        # Use thread connection for parallel execution
                        df = analyzer.execute_sql(sql_stmt, use_thread_connection=True)
                        return df, idx, item
                    except Exception as e:
                        print(f"❌ Error executing query {idx}: {e}")
                        return None, idx, item
                
                # Execute queries in parallel (max 5 concurrent)
                with ThreadPoolExecutor(max_workers=5) as executor:
                    future_to_idx = {
                        executor.submit(execute_single_query, idx, item): idx 
                        for idx, item in enumerate(sql_items, start=1)
                    }
                    completed_results = {}
                    for future in as_completed(future_to_idx):
                        df, idx, item = future.result()
                        completed_results[idx] = (df, item)
                
                # Process results in original order and stream progress
                for idx in sorted(completed_results.keys()):
                    df, item = completed_results[idx]
                    sql_stmt = (item.get('sql') or '').strip().rstrip(';')
                    rows = 0
                    cols = []
                    if df is not None:
                        try:
                            rows = int(df.shape[0])
                            cols = list(df.columns)
                            if rows == 0:
                                print(f"⚠️ Query {idx} returned empty result")
                            else:
                                print(f"✅ Query {idx} returned {rows} rows")
                        except Exception as e:
                            print(f"⚠️ Error processing DataFrame for query {idx}: {e}")
                            rows = 0
                            cols = []
                        # Add dataset even if empty (for analysis to report "no data")
                        datasets.append({
                            'name': item.get('name'),
                            'purpose': item.get('purpose'),
                            'sql': sql_stmt,
                            'df': df
                        })
                    else:
                        print(f"❌ Query {idx} execution failed (returned None)")
                    yield sse('execution_progress', {
                        'index': idx,
                        'name': item.get('name') or f'Dataset {idx}',
                        'rows': rows,
                        'columns': cols
                    })

                # Check if we have any datasets at all
                if not datasets:
                    print("❌ No valid datasets generated - all queries failed")
                    yield sse('analysis_final', {'text': 'No data found for the generated queries. All SQL queries either failed or returned no results.'})
                    yield sse('done', {'error': 'No datasets'})
                    return
                
                # Check if all datasets are empty
                all_empty = True
                for d in datasets:
                    if d['df'] is not None:
                        try:
                            if not d['df'].empty:
                                all_empty = False
                                break
                        except Exception:
                            pass
                
                if all_empty:
                    print("⚠️ All queries returned empty results")
                    # Still proceed with analysis so AI can report "no data found"

                # Stream analysis tokens (same as initial query)
                accumulated = ''
                for token in analyzer.analyze_multi_data(datasets, enhanced_question, mode=selected_option, stream=True, session_id=session_id):
                    if token:
                        accumulated += token
                        yield sse('analysis_token', {'text': token})
                yield sse('analysis_final', {'text': accumulated})

                # Add to existing conversation as followup
                try:
                    # Get current conversation context
                    context = analyzer.conversation_manager.get_conversation_details(session_id)
                    if not context:
                        print(f"⚠️ Warning: Session {session_id} not found, cannot save followup")
                    else:
                        # Choose primary df compatible with existing storage
                        primary_df = None
                        for d in datasets:
                            if d["df"] is not None and not d["df"].empty:
                                primary_df = d["df"]
                                break
                        if primary_df is None:
                            primary_df = datasets[0]["df"] if datasets else None

                        # Join SQLs for storage
                        joined_sql = "\n\n".join([f"-- {i+1}. {item.get('name','Query')}\n{item.get('sql','')}" for i, item in enumerate(sql_items)])
                        
                        # Add followup with complete analysis data
                        followup_data = {
                            "timestamp": datetime.now().isoformat(),
                            "question": question,
                            "sql_query": joined_sql,
                            "query_result": primary_df.to_dict('records') if primary_df is not None and hasattr(primary_df, 'to_dict') else [],
                            "columns": list(primary_df.columns) if primary_df is not None and hasattr(primary_df, 'columns') else [],
                            "analysis": accumulated,
                            "type": "followup",
                            "selected_option": selected_option
                        }
                        if context.conversation_history is None:
                            context.conversation_history = []
                        context.conversation_history.append(followup_data)
                        analyzer.conversation_manager.session_manager.update_session(session_id, context)
                        print(f"✅ Successfully saved followup to conversation session_id={session_id}")
                except Exception as e:
                    print(f"❌ Error saving followup to database: {e}")
                    print(f"   Session ID: {session_id}")
                    traceback.print_exc()
                    # Don't fail the entire request, but log the error

                # Auto-generate DOC for DataReports
                doc_file_url = None
                doc_filename = None
                if selected_option == 'DataReports' and template_id:
                    try:
                        doc_result = auto_generate_doc_for_data_reports(session_id, question, accumulated, template_id)
                        if 'error' not in doc_result:
                            doc_file_url = doc_result.get('doc_file_url')
                            doc_filename = doc_result.get('filename')
                            # Persist DOC info to conversation
                            try:
                                context = analyzer.conversation_manager.get_conversation_details(session_id)
                                if context:
                                    context.doc_file_url = doc_file_url
                                    context.doc_filename = doc_filename
                                    analyzer.conversation_manager.session_manager.update_session(session_id, context)
                            except Exception:
                                pass
                    except Exception:
                        pass

                # Handle Charts mode - generate chart for followup questions
                has_chart = False
                chart_spec = None
                chart_data = None
                chart_error = None
                
                if selected_option == 'Charts' and chart_type and primary_df is not None:
                    try:
                        # Convert query data to list of dictionaries if needed
                        if hasattr(primary_df, 'to_dict'):
                            chart_data = primary_df.to_dict('records')
                        elif isinstance(primary_df, list):
                            chart_data = primary_df
                        else:
                            chart_data = None
                        
                        if chart_data and len(chart_data) > 0:
                            # Generate chart using chart generator
                            chart_spec = chart_generator.generate_chart(
                                data=chart_data,
                                chart_type=chart_type
                            )
                            
                            if chart_spec:
                                has_chart = True
                                print(f"Chart generated successfully for followup stream with type: {chart_type}")
                            else:
                                chart_error = "Failed to generate chart specification"
                                print("Failed to generate chart specification for followup stream")
                        else:
                            chart_error = "No data available for chart generation"
                            print("No data available for chart generation in followup stream")
                            
                    except Exception as e:
                        print(f"Error generating chart for followup stream: {e}")
                        chart_error = f"Chart generation error: {str(e)}"

                # Handle Excel mode - generate Excel file for followup questions
                excel_file_url = None
                excel_filename = None
                auto_excel_generated = False

                if selected_option == 'Excel' and excel_type and primary_df is not None:
                    try:
                        # Generate Excel file using correct signature
                        excel_path, excel_name = create_excel_file(
                            excel_type=excel_type,
                            data=primary_df if primary_df is not None else [],
                            analysis_text=accumulated,
                            session_id=session_id
                        )

                        # Create download URL and set flags
                        excel_file_url = f"/api/download/{excel_name}"
                        excel_filename = excel_name
                        auto_excel_generated = True
                        print(f"Excel file generated successfully for followup stream with type: {excel_type} -> {excel_name}")

                        # Persist Excel info to conversation
                        try:
                            context = analyzer.conversation_manager.get_conversation_details(session_id)
                            if context:
                                context.excel_file_url = excel_file_url
                                context.excel_filename = excel_filename
                                context.auto_excel_generated = auto_excel_generated
                                analyzer.conversation_manager.session_manager.update_session(session_id, context)
                        except Exception as e:
                            print(f"Error persisting Excel info to conversation: {e}")
                    except Exception as e:
                        print(f"Error generating Excel file for followup stream: {e}")

                yield sse('done', {
                    'session_id': session_id,
                    'selected_option': selected_option,
                    'template_id': template_id,
                    'chart_type': chart_type,
                    'excel_type': excel_type,
                    'doc_file_url': doc_file_url,
                    'doc_filename': doc_filename,
                    'has_chart': has_chart,
                    'chart_spec': chart_spec,
                    'chart_data': chart_data,
                    'chart_error': chart_error,
                    'excel_file_url': excel_file_url,
                    'excel_filename': excel_filename,
                    'auto_excel_generated': auto_excel_generated
                })
            except Exception as e:
                yield sse('done', {'error': str(e)})

        resp = Response(generate(), mimetype='text/event-stream')
        resp.headers['Cache-Control'] = 'no-cache'
        resp.headers['X-Accel-Buffering'] = 'no'
        resp.headers['Connection'] = 'keep-alive'
        resp.headers['Access-Control-Allow-Origin'] = '*'
        resp.headers['Access-Control-Allow-Headers'] = 'Content-Type'
        return resp
    except Exception as e:
        print(f"Error in /api/followup/stream: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/followup', methods=['POST'])
def ask_followup_question():
    """Ask a followup question in an existing conversation - now uses complete data analysis workflow"""
    try:
        data = request.get_json()
        session_id = data.get('session_id')
        question = data.get('question')
        
        if not session_id or not question:
            return jsonify({'error': 'Session ID and question are required'}), 400

        # CRITICAL: Reject temporary session IDs to prevent data loss
        if isinstance(session_id, str) and (session_id.startswith('temp_') or session_id.startswith('temp-')):
            return jsonify({
                'error': 'Invalid session ID. The previous conversation was not saved. Please start a new conversation.',
                'session_id': None
            }), 400
        
        # Extract template information from question if present
        template_id = None
        enhanced_question = question
        selected_option = data.get('selected_option')  # Get from request data
        excel_type = data.get('excel_type')  # Get Excel type from request data
        
        # Check for different option patterns
        data_reports_pattern = r'\[DataReports - (\w+)\]'
        charts_pattern = r'\[Charts - (\w+)\]'
        excel_pattern = r'\[Excel - (\w+)\]'
        general_pattern = r'\[(\w+)\]'
        
        # Check if question contains specific template markers
        data_reports_match = re.search(data_reports_pattern, question)
        charts_match = re.search(charts_pattern, question)
        excel_match = re.search(excel_pattern, question)
        general_match = re.search(general_pattern, question)
        
        chart_type = None
        
        if data_reports_match:
            template_id = data_reports_match.group(1)
            selected_option = 'DataReports'
            # Remove the marker from the question
            clean_question = re.sub(data_reports_pattern, '', question).strip()
            # Enhance the question with template-specific prompt
            enhanced_question = enhance_prompt_with_template(template_id, clean_question)
            print(f"Using template in followup: {template_id}")
            print(f"Enhanced followup question: {enhanced_question}")
        elif charts_match:
            chart_type = charts_match.group(1)
            selected_option = 'Charts'
            enhanced_question = re.sub(charts_pattern, '', question).strip()
        elif excel_match:
            excel_type = excel_match.group(1)
            selected_option = 'Excel'
            clean_question = re.sub(excel_pattern, '', question).strip()
            enhanced_question = enhance_prompt_with_excel_template(excel_type, clean_question)
        elif general_match:
            selected_option = general_match.group(1)
            # Remove the marker from the question
            enhanced_question = re.sub(general_pattern, '', question).strip()
            print(f"Using {selected_option} mode in followup")
        
        # Check if user is requesting DOC generation
        doc_keywords = ['yes', 'generate doc', 'create doc', 'make doc', 'doc please', 'generate document', 'create document']
        user_wants_doc = any(keyword in enhanced_question.lower() for keyword in doc_keywords)
        
        if user_wants_doc and DOC_EXPORT_AVAILABLE:
            # Generate DOC file automatically
            try:
                # Get conversation details for DOC generation
                conversation_details = analyzer.conversation_manager.get_conversation_details(session_id)
                if conversation_details:
                    # Create DOC file
                    doc = Document()
                    
                    # Add title
                    title = doc.add_heading('Data Analysis Report', 0)
                    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    # Add metadata
                    doc.add_heading('Report Information', level=1)
                    info_table = doc.add_table(rows=4, cols=2)
                    info_table.style = 'Table Grid'
                    
                    info_table.cell(0, 0).text = 'Session ID'
                    info_table.cell(0, 1).text = str(conversation_details.session_id)
                    info_table.cell(1, 0).text = 'Generated Date'
                    info_table.cell(1, 1).text = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                    info_table.cell(2, 0).text = 'Language'
                    info_table.cell(2, 1).text = conversation_details.language or 'English'
                    info_table.cell(3, 0).text = 'Template Used'
                    info_table.cell(3, 1).text = template_id or 'None'
                    
                    # Add conversation content
                    doc.add_heading('Analysis Content', level=1)
                    
                    # Add original question
                    doc.add_heading('Original Question', level=2)
                    doc.add_paragraph(conversation_details.original_question)
                    
                    # Add analysis
                    doc.add_heading('Analysis Results', level=2)
                    doc.add_paragraph(conversation_details.analysis_result)
                    
                    # Add conversation history if available
                    if conversation_details.conversation_history:
                        doc.add_heading('Conversation History', level=2)
                        try:
                            import json
                            history = json.loads(conversation_details.conversation_history)
                            for i, entry in enumerate(history, 1):
                                if isinstance(entry, dict):
                                    doc.add_heading(f'Exchange {i}', level=3)
                                    if 'question' in entry:
                                        doc.add_paragraph(f"Q: {entry['question']}")
                                    if 'analysis' in entry:
                                        doc.add_paragraph(f"A: {entry['analysis']}")
                        except:
                            doc.add_paragraph(conversation_details.conversation_history)
                    
                    # Save to temporary file
                    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
                    doc.save(temp_file.name)
                    temp_file.close()
                    
                    # Generate download URL (simplified - in production you'd want a proper file serving mechanism)
                    import uuid
                    file_id = str(uuid.uuid4())[:8]
                    filename = f"analysis_report_{file_id}.docx"
                    
                    # Move file to exports directory
                    import shutil
                    export_path = os.path.join(os.path.dirname(__file__), 'exports', filename)
                    shutil.move(temp_file.name, export_path)
                    
                    # Create download link response
                    download_link = f"http://localhost:5000/api/download/{filename}"
                    doc_response = f"✅ **DOC Report Generated Successfully!**\n\n📄 Your comprehensive analysis report has been created and is ready for download.\n\n🔗 **[Click here to download your report]({download_link})**\n\nThe document includes:\n- Complete analysis results\n- Professional formatting\n- Template-specific content\n- Conversation history\n\n*File: {filename}*"
                    
                    response = {
                        'session_id': session_id,
                        'question': question,
                        'analysis': doc_response,
                        'timestamp': datetime.now().isoformat(),
                        'template_id': template_id,
                        'template_config': get_template_by_sub_option(template_id) if template_id else None,
                        'selected_option': selected_option,
                        'doc_generated': True,
                        'download_link': download_link,
                        'filename': filename
                    }
                    
                    return jsonify(response)
                    
            except Exception as e:
                print(f"Error generating DOC: {e}")
                # Fall back to normal response with error message
                error_response = f"❌ **Sorry, I encountered an error while generating the DOC file.**\n\nError: {str(e)}\n\nPlease try again or contact support if the issue persists."
                response = {
                    'session_id': session_id,
                    'question': question,
                    'analysis': error_response,
                    'timestamp': datetime.now().isoformat(),
                    'template_id': template_id,
                    'template_config': get_template_by_sub_option(template_id) if template_id else None,
                    'selected_option': selected_option,
                    'doc_generation_failed': True
                }
                return jsonify(response)
        
        # Use complete data analysis workflow for followup (same as initial query)
        # Get language from existing conversation context
        context = analyzer.conversation_manager.get_conversation_details(session_id)
        language = context.language if context else 'chinese'
        
        # Use the LLM-first multi-SQL workflow as initial queries
        query_data, analysis, _ = ask_and_analyze_v2(
            analyzer,
            enhanced_question, 
            create_session=False,  # Don't create new session, add to existing one
            language=language,
            selected_option=selected_option,
            chart_type=chart_type,
            excel_type=excel_type,
            session_id=session_id  # Pass session_id to add to existing conversation
        )
        
        # Manually save follow-up to existing conversation
        if session_id and query_data is not None:
            try:
                context = analyzer.conversation_manager.get_conversation_details(session_id)
                if context:
                    # Prepare follow-up data
                    primary_df = query_data
                    joined_sql = "\n\n".join([
                        f"-- Query {i+1}\n{sql}" 
                        for i, sql in enumerate(analyzer.last_sql_list or [])
                    ]) if hasattr(analyzer, 'last_sql_list') and analyzer.last_sql_list else ""
                    
                    followup_data = {
                        'timestamp': datetime.now().isoformat(),
                        'question': question,
                        'sql_query': joined_sql,
                        'query_result': primary_df.to_dict('records') if hasattr(primary_df, 'to_dict') else [],
                        'columns': list(primary_df.columns) if hasattr(primary_df, 'columns') else [],
                        'analysis': analysis,
                        'type': 'followup',
                        'selected_option': selected_option
                    }
                    if context.conversation_history is None:
                        context.conversation_history = []
                    context.conversation_history.append(followup_data)
                    analyzer.conversation_manager.session_manager.update_session(session_id, context)
                    print(f"✅ Successfully saved followup to conversation session_id={session_id}")
                else:
                    print(f"⚠️ Warning: Session {session_id} not found, cannot save followup")
            except Exception as e:
                print(f"❌ Error saving followup to database: {e}")
                print(f"   Session ID: {session_id}")
                traceback.print_exc()
        
        response = {
            'session_id': session_id,
            'question': question,
            'analysis': analysis,
            'query_data': query_data,
            'timestamp': datetime.now().isoformat(),
            'template_id': template_id,
            'template_config': get_template_by_sub_option(template_id) if template_id else None,
            'selected_option': selected_option
        }
        
        # Handle Charts mode - generate chart for followup questions
        if selected_option == 'Charts' and chart_type and query_data is not None:
            try:
                # Convert query data to list of dictionaries if needed
                if hasattr(query_data, 'to_dict'):
                    chart_data = query_data.to_dict('records')
                elif isinstance(query_data, list):
                    chart_data = query_data
                else:
                    chart_data = None
                
                if chart_data and len(chart_data) > 0:
                    # Generate chart using chart generator
                    chart_spec = chart_generator.generate_chart(
                        data=chart_data,
                        chart_type=chart_type
                    )
                    
                    if chart_spec:
                        response['has_chart'] = True
                        response['chart_spec'] = chart_spec
                        response['chart_data'] = chart_data
                        response['chart_type'] = chart_type
                        print(f"Chart generated successfully for followup question with type: {chart_type}")
                    else:
                        response['has_chart'] = False
                        response['chart_error'] = "Failed to generate chart specification"
                        print("Failed to generate chart specification for followup")
                else:
                    response['has_chart'] = False
                    response['chart_error'] = "No data available for chart generation"
                    print("No data available for chart generation in followup")
                    
            except Exception as e:
                print(f"Error generating chart for followup: {e}")
                response['has_chart'] = False
                response['chart_error'] = f"Chart generation error: {str(e)}"
        
        # Auto-generate DOC file for DataReports mode
        if selected_option == 'DataReports':
            # Automatically generate DOC file
            doc_result = auto_generate_doc_for_data_reports(session_id, question, analysis, template_id)
            
            if 'error' in doc_result:
                # If DOC generation failed, add error message to analysis
                doc_error_msg = f"\n\n---\n\n**DOC Generation Error:** {doc_result['error']}"
                response['analysis'] = analysis + doc_error_msg
                response['doc_generation_error'] = doc_result['error']
            else:
                # If DOC generation succeeded, add download link to analysis
                doc_success_msg = f"\n\n---\n\n**📄 DOC Report Generated Successfully!**\n\nYour comprehensive data analysis report has been automatically generated and is ready for download.\n\n**File:** {doc_result['filename']}\n\n[Download DOC Report]({doc_result['doc_file_url']})"
                response['analysis'] = analysis + doc_success_msg
                response['doc_file_url'] = doc_result['doc_file_url']
                response['doc_filename'] = doc_result['filename']
                response['auto_doc_generated'] = True
                # Persist DOC info to conversation context
                try:
                    context = analyzer.conversation_manager.get_conversation_details(session_id)
                    if context:
                        context.doc_file_url = response['doc_file_url']
                        context.doc_filename = response['doc_filename']
                        analyzer.conversation_manager.session_manager.update_session(session_id, context)
                except Exception:
                    pass
        
        # Auto-generate Excel file for Excel mode
        if selected_option == 'Excel' and excel_type:
            try:
                excel_path, excel_name = create_excel_file(
                    excel_type=excel_type,
                    data=query_data if query_data is not None else [],
                    analysis_text=analysis,
                    session_id=session_id
                )
                excel_file_url = f"/api/download/{excel_name}"
                excel_filename = excel_name
                
                # Add success message to analysis
                excel_success_msg = f"\n\n---\n\n**📊 Excel Report Generated Successfully!**\n\nYour {excel_type} analysis Excel file has been automatically generated and is ready for download.\n\n**File:** {excel_filename}\n\n[Download Excel Report]({excel_file_url})"
                response['analysis'] = analysis + excel_success_msg
                response['excel_file_url'] = excel_file_url
                response['excel_filename'] = excel_filename
                response['auto_excel_generated'] = True
                
                # Persist Excel info to conversation context
                try:
                    context = analyzer.conversation_manager.get_conversation_details(session_id)
                    if context:
                        context.excel_file_url = excel_file_url
                        context.excel_filename = excel_filename
                        analyzer.conversation_manager.session_manager.update_session(session_id, context)
                except Exception:
                    pass
            except Exception as e:
                print(f"Error generating Excel for followup: {e}")
                excel_error_msg = f"\n\n---\n\n**Excel Generation Error:** {str(e)}"
                response['analysis'] = analysis + excel_error_msg
                response['excel_generation_error'] = str(e)
        
        return jsonify(response)
        
    except Exception as e:
        print(f"Error processing followup question: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/export/doc', methods=['POST'])
def export_to_doc():
    """Export conversation to DOC format"""
    try:
        if not DOC_EXPORT_AVAILABLE:
            return jsonify({'error': 'DOC export not available. Please install python-docx.'}), 500
            
        data = request.get_json()
        session_id = data.get('session_id')
        
        if not session_id:
            return jsonify({'error': 'Session ID is required'}), 400
        
        # Get conversation details
        conversation = analyzer.get_session_details(session_id)
        if not conversation:
            return jsonify({'error': 'Conversation not found'}), 404
        
        # Create DOC document
        doc = Document()
        
        # Add title
        title = doc.add_heading('Craveva AI Enterprise Data Analysis Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add metadata
        doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"Session ID: {session_id}")
        doc.add_paragraph("")
        
        # Add conversation content
        for i, qa in enumerate(conversation.get('conversation', []), 1):
            # Add question
            question_heading = doc.add_heading(f"Question {i}", level=2)
            doc.add_paragraph(qa.get('question', ''))
            
            # Add analysis
            if qa.get('analysis'):
                doc.add_heading('Analysis:', level=3)
                doc.add_paragraph(qa.get('analysis', ''))
            
            # Add data table if available
            if qa.get('data') and len(qa.get('data', [])) > 0:
                doc.add_heading('Data Results:', level=3)
                
                # Create table
                data_list = qa.get('data', [])
                if data_list:
                    columns = qa.get('columns', list(data_list[0].keys()) if data_list else [])
                    
                    # Add table
                    table = doc.add_table(rows=1, cols=len(columns))
                    table.style = 'Table Grid'
                    
                    # Add header row
                    hdr_cells = table.rows[0].cells
                    for i, column in enumerate(columns):
                        hdr_cells[i].text = str(column)
                    
                    # Add data rows (limit to first 50 rows for performance)
                    for row_data in data_list[:50]:
                        row_cells = table.add_row().cells
                        for i, column in enumerate(columns):
                            row_cells[i].text = str(row_data.get(column, ''))
                    
                    if len(data_list) > 50:
                        doc.add_paragraph(f"Note: Showing first 50 rows of {len(data_list)} total rows.")
            
            # Add template information if available
            if qa.get('template_id'):
                doc.add_heading('Template Used:', level=3)
                doc.add_paragraph(f"Template ID: {qa.get('template_id')}")
                template_config = qa.get('template_config')
                if template_config:
                    doc.add_paragraph(f"Template Name: {template_config.get('name', 'N/A')}")
                    doc.add_paragraph(f"Description: {template_config.get('description', 'N/A')}")
            
            doc.add_paragraph("")  # Add spacing
        
        # Save to temporary file
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
        doc.save(temp_file.name)
        temp_file.close()
        
        # Generate filename
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"coffee_analysis_report_{timestamp}.docx"
        
        return send_file(
            temp_file.name,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        
    except Exception as e:
        print(f"Error exporting to DOC: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/download/<filename>', methods=['GET'])
def download_file(filename):
    """Download generated files"""
    try:
        # Security check - only allow downloading from exports directory
        if '..' in filename or '/' in filename or '\\' in filename:
            return jsonify({'error': 'Invalid filename'}), 400
        
        export_path = os.path.join(os.path.dirname(__file__), 'exports', filename)
        
        if not os.path.exists(export_path):
            return jsonify({'error': 'File not found'}), 404
        
        return send_file(export_path, as_attachment=True, download_name=filename)
        
    except Exception as e:
        print(f"Error downloading file: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/export', methods=['POST'])
def export_conversation():
    """Export conversation to PDF or CSV"""
    try:
        data = request.get_json()
        session_id = data.get('session_id')
        format_type = data.get('format', 'pdf')
        
        if not session_id:
            return jsonify({'error': 'Session ID is required'}), 400
        
        if format_type == 'pdf':
            file_path = analyzer.export_session_to_pdf(session_id)
        elif format_type == 'csv':
            file_path = analyzer.export_session_to_csv(session_id)
        else:
            return jsonify({'error': 'Invalid format. Use pdf or csv'}), 400
        
        if file_path:
            return jsonify({
                'success': True,
                'file_path': file_path,
                'message': f'Conversation exported to {format_type.upper()}'
            })
        else:
            return jsonify({'error': 'Export failed'}), 500
            
    except Exception as e:
        print(f"Error exporting conversation: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/export/response', methods=['POST'])
def export_single_response():
    """Export single AI response to PDF without user questions and SQL queries"""
    try:
        data = request.get_json()
        session_id = data.get('session_id')
        analysis_text = data.get('analysis_text')
        
        if not session_id:
            return jsonify({'error': 'Session ID is required'}), 400
        
        if not analysis_text:
            return jsonify({'error': 'Analysis text is required'}), 400
        
        file_path = analyzer.export_single_ai_response_to_pdf(session_id, analysis_text)
        
        if file_path:
            return jsonify({
                'success': True,
                'file_path': file_path,
                'message': 'AI response exported to PDF'
            })
        else:
            return jsonify({'error': 'Export failed'}), 500
            
    except Exception as e:
        print(f"Error exporting single response: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/conversations/<session_id>', methods=['DELETE'])
def delete_conversation(session_id):
    """Delete a specific conversation"""
    try:
        if not session_id:
            return jsonify({'error': 'Session ID is required'}), 400
        
        # Use the analyzer to delete the conversation
        success = analyzer.conversation_manager.delete_conversation(session_id)
        
        if success:
            return jsonify({
                'success': True,
                'message': 'Conversation deleted successfully'
            })
        else:
            return jsonify({'error': 'Conversation not found or deletion failed'}), 404
            
    except Exception as e:
        print(f"Error deleting conversation: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/export/ppt', methods=['POST'])
def export_to_ppt():
    """Export analysis to PowerPoint presentation"""
    try:
        if not PPT_EXPORT_AVAILABLE:
            return jsonify({'error': 'PPT export not available. Please install python-pptx.'}), 500
        
        data = request.get_json()
        session_id = data.get('session_id')
        analysis_text = data.get('analysis_text', '')
        ppt_type = data.get('ppt_type', 'Proposal')  # Default to Proposal
        question = data.get('question', '')
        
        if not session_id:
            return jsonify({'error': 'Session ID is required'}), 400
        
        # Map frontend button names to PPT generator function names
        ppt_type_mapping = {
            'Products & Services': 'ProductsServices',
            'Market & Clients': 'MarketClients', 
            'Performance & Growth': 'PerformanceGrowth',
            'Strategy & Vision': 'StrategyVision',
            'Proposal': 'Proposal'
        }
        
        # Get the mapped type or use the original if not found
        mapped_type = ppt_type_mapping.get(ppt_type, ppt_type)
        
        # Generate PPT using the appropriate template
        file_path = ppt_generator.generate_ppt_by_type(
            ppt_type=mapped_type,
            analysis_data=analysis_text,
            conversation_id=session_id
        )
        
        if file_path and os.path.exists(file_path):
            # Get just the filename for the response
            filename = os.path.basename(file_path)
            
            return jsonify({
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'download_url': f'/api/download/{filename}',
                'message': f'PPT presentation generated successfully for {ppt_type}'
            })
        else:
            return jsonify({'error': 'PPT generation failed'}), 500
            
    except Exception as e:
        print(f"Error generating PPT: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/ppt/generate', methods=['POST'])
def generate_ppt_from_conversation():
    """Generate PPT from conversation with AI dialogue support"""
    try:
        if not PPT_EXPORT_AVAILABLE:
            return jsonify({'error': 'PPT export not available. Please install python-pptx.'}), 500
        
        data = request.get_json()
        question = data.get('question')
        ppt_type = data.get('ppt_type', 'Proposal')
        session_id = data.get('session_id')
        
        if not question:
            return jsonify({'error': 'Question is required'}), 400
        
        # If no session_id provided, create a new one
        if not session_id:
            session_id = str(uuid.uuid4())
        
        # Process the question with AI to get analysis
        enhanced_question = f"[PPT - {ppt_type}] {question}"
        
        # Get AI analysis
        try:
            # Try the main analyzer first (if available)
            if hasattr(analyzer, 'ask_and_analyze') and 'session_id' in analyzer.ask_and_analyze.__code__.co_varnames:
                query_data, analysis, session_id = analyzer.ask_and_analyze(
                    enhanced_question, 
                    session_id=session_id,
                    selected_option='PPT'
                )
            else:
                # Fallback to SimpleCoffeeAnalyzer
                result = analyzer.ask_and_analyze(enhanced_question)
                # Convert DataFrame to string if needed
                if hasattr(result, 'to_string'):
                    analysis = result.to_string()
                else:
                    analysis = str(result)
                query_data = None
        except Exception as e:
            # Fallback to SimpleCoffeeAnalyzer
            result = analyzer.ask_and_analyze(enhanced_question)
            # Convert DataFrame to string if needed
            if hasattr(result, 'to_string'):
                analysis = result.to_string()
            else:
                analysis = str(result)
            query_data = None
        
        # Map frontend button names to PPT generator function names
        ppt_type_mapping = {
            'Products & Services': 'ProductsServices',
            'Market & Clients': 'MarketClients', 
            'Performance & Growth': 'PerformanceGrowth',
            'Strategy & Vision': 'StrategyVision',
            'Proposal': 'Proposal'
        }
        
        # Get the mapped type or use the original if not found
        mapped_type = ppt_type_mapping.get(ppt_type, ppt_type)
        
        # Generate PPT using the analysis
        file_path = ppt_generator.generate_ppt_by_type(
            ppt_type=mapped_type,
            analysis_data=analysis,
            conversation_id=session_id
        )
        
        if file_path and os.path.exists(file_path):
            # Get just the filename for the response
            filename = os.path.basename(file_path)
            
            return jsonify({
                'success': True,
                'session_id': session_id,
                'analysis': analysis,
                'file_path': file_path,
                'filename': filename,
                'download_url': f'/api/download/{filename}',
                'message': f'PPT presentation generated successfully for {ppt_type}',
                'ppt_type': ppt_type
            })
        else:
            return jsonify({
                'success': False,
                'session_id': session_id,
                'analysis': analysis,
                'error': 'PPT generation failed',
                'ppt_type': ppt_type
            }), 500
            
    except Exception as e:
        print(f"Error generating PPT from conversation: {e}")
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/health', methods=['GET'])
def health_check():
    """Health check endpoint"""
    return jsonify({
        'status': 'healthy',
        'timestamp': datetime.now().isoformat(),
        'version': '1.0.0'
    })

@app.route('/api/vector-db/status', methods=['GET'])
def get_vector_db_status():
    """Get vector database status"""
    try:
        status = analyzer.get_vector_db_status()
        return jsonify(status)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/vector-db/refresh', methods=['POST'])
def refresh_vector_db():
    """Refresh/reload vector database schema from MySQL"""
    try:
        result = analyzer.refresh_vector_db_schema()
        if result['status'] == 'success':
            return jsonify(result), 200
        else:
            return jsonify(result), 500
    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)}), 500

@app.errorhandler(404)
def not_found(error):
    return jsonify({'error': 'Endpoint not found'}), 404

@app.errorhandler(500)
def internal_error(error):
    return jsonify({'error': 'Internal server error'}), 500

if __name__ == '__main__':
    print("Starting Craveva AI Enterprise Business Web API...")
    print("Frontend should be accessible at: http://localhost:3000")
    print("API endpoints available at: http://localhost:5000/api/")
    print("\nAvailable endpoints:")
    print("  GET    /api/health - Health check")
    print("  GET    /api/conversations - List all conversations")
    print("  GET    /api/conversations/<id> - Get conversation details")
    print("  DELETE /api/conversations/<id> - Delete a conversation")
    print("  POST   /api/ask - Ask a new question")
    print("  POST   /api/followup - Ask a followup question")
    print("  POST   /api/export - Export conversation")
    
    port = int(os.environ.get('PORT', 5000))
    
    # Add signal handlers for graceful shutdown
    import signal
    import sys
    
    def signal_handler(sig, frame):
        print("\n⚠️ Received shutdown signal, shutting down gracefully...")
        sys.exit(0)
    
    signal.signal(signal.SIGINT, signal_handler)
    signal.signal(signal.SIGTERM, signal_handler)
    
    try:
        app.run(debug=False, use_reloader=False, host='0.0.0.0', port=port, threaded=True)
    except Exception as e:
        print(f"❌ Failed to start server: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)