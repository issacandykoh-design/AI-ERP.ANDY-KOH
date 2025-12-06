"""
Excel Generator Module
Generates Excel files based on different business categories and data analysis results.
"""

import os
import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd
import re


class ExcelGenerator:
    def __init__(self, exports_dir="exports"):
        """Initialize Excel generator with exports directory."""
        self.exports_dir = exports_dir
        if not os.path.exists(exports_dir):
            os.makedirs(exports_dir)
    
    def create_workbook(self, title="Data Analysis"):
        """Create a new workbook with basic styling."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Analysis"
        
        # Set default styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        return wb, ws, header_font, header_fill
    
    def apply_header_style(self, ws, row_num, header_font, header_fill):
        """Apply header styling to a specific row."""
        for cell in ws[row_num]:
            if cell.value:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
    
    def generate_sales_excel(self, data, analysis_text, session_id):
        """Generate Sales-focused Excel file."""
        wb, ws, header_font, header_fill = self.create_workbook("Sales Analysis")
        
        # Title
        ws['A1'] = "Sales Analysis Report"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:E1')
        
        # Date
        ws['A2'] = f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws.merge_cells('A2:E2')
        
        # Analysis Summary
        ws['A4'] = "Analysis Summary"
        ws['A4'].font = Font(size=14, bold=True)
        ws['A5'] = analysis_text if analysis_text else "No analysis available"
        # Dynamically adjust merge range based on content length
        merge_rows = max(8, min(20, len(analysis_text) // 100)) if analysis_text else 8
        ws.merge_cells(f'A5:E{merge_rows}')
        ws['A5'].alignment = Alignment(wrap_text=True, vertical="top")
        
        # Data section starts after analysis
        data_start_row = merge_rows + 3
        
        # Check if we have actual DataFrame data
        if data is not None and hasattr(data, 'columns') and not data.empty:
            # Add actual data from DataFrame
            ws.cell(row=data_start_row, column=1, value="Raw Sales Data")
            ws.cell(row=data_start_row, column=1).font = Font(size=14, bold=True)
            
            # Add DataFrame headers
            data_header_row = data_start_row + 2
            for col_idx, column_name in enumerate(data.columns, 1):
                ws.cell(row=data_header_row, column=col_idx, value=str(column_name))
            
            self.apply_header_style(ws, data_header_row, header_font, header_fill)
            
            # Add DataFrame data (all rows, no limit)
            for row_idx, (_, row_data) in enumerate(data.iterrows(), data_header_row + 1):
                for col_idx, value in enumerate(row_data, 1):
                    # Convert value to string to handle different data types
                    cell_value = str(value) if pd.notna(value) else ""
                    ws.cell(row=row_idx, column=col_idx, value=cell_value)
            
            # Add sales summary after the data
            summary_start_row = data_header_row + len(data) + 3
            ws.cell(row=summary_start_row, column=1, value="Sales Summary")
            ws.cell(row=summary_start_row, column=1).font = Font(size=14, bold=True)
            
            # Calculate and add summary statistics
            summary_row = summary_start_row + 2
            if 'money' in data.columns:
                total_sales = data['money'].sum()
                transaction_count = len(data)
                avg_sale = data['money'].mean()
                
                summary_data = [
                    ["Total Sales", f"¥{total_sales:,.2f}"],
                    ["Total Transactions", f"{transaction_count:,}"],
                    ["Average Sale", f"¥{avg_sale:.2f}"]
                ]
                
                for i, (label, value) in enumerate(summary_data, summary_row):
                    ws.cell(row=i, column=1, value=label)
                    ws.cell(row=i, column=2, value=value)
                    ws.cell(row=i, column=1).font = Font(bold=True)
        else:
            # Fallback to sample data if no actual data provided
            headers = ["Product", "Sales Amount", "Quantity", "Revenue", "Growth %"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=data_start_row, column=col, value=header)
            
            self.apply_header_style(ws, data_start_row, header_font, header_fill)
            
            sample_data = [
                ["Product A", 15000, 120, 18000, "12%"],
                ["Product B", 22000, 180, 26400, "8%"],
                ["Product C", 18500, 150, 22200, "15%"],
                ["Total", 55500, 450, 66600, "11.7%"]
            ]
            
            for i, row_data in enumerate(sample_data, data_start_row + 1):
                for col, value in enumerate(row_data, 1):
                    ws.cell(row=i, column=col, value=value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column:
                try:
                    # Skip merged cells
                    if hasattr(cell, 'column_letter'):
                        if column_letter is None:
                            column_letter = cell.column_letter
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                except:
                    pass
            if column_letter:
                # Remove width limit to allow full content display
                adjusted_width = max_length + 2
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save file
        filename = f"excel_Sales_Analysis_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.exports_dir, filename)
        wb.save(filepath)
        
        return filepath, filename
    
    def generate_finance_excel(self, data, analysis_text, session_id):
        """Generate Finance-focused Excel file."""
        wb, ws, header_font, header_fill = self.create_workbook("Finance Analysis")
        
        # Title
        ws['A1'] = "Financial Analysis Report"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:F1')
        
        # Date
        ws['A2'] = f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws.merge_cells('A2:F2')
        
        # Analysis Summary
        ws['A4'] = "Financial Analysis Summary"
        ws['A4'].font = Font(size=14, bold=True)
        ws['A5'] = analysis_text if analysis_text else "No analysis available"
        # Dynamically adjust merge range based on content length
        merge_rows = max(8, min(20, len(analysis_text) // 100)) if analysis_text else 8
        ws.merge_cells(f'A5:F{merge_rows}')
        ws['A5'].alignment = Alignment(wrap_text=True, vertical="top")
        
        # Data section starts after analysis
        data_start_row = merge_rows + 3
        
        # Check if we have actual DataFrame data
        if data is not None and hasattr(data, 'columns') and not data.empty:
            # Add actual data from DataFrame
            ws.cell(row=data_start_row, column=1, value="Raw Transaction Data")
            ws.cell(row=data_start_row, column=1).font = Font(size=14, bold=True)
            
            # Add DataFrame headers
            data_header_row = data_start_row + 2
            for col_idx, column_name in enumerate(data.columns, 1):
                ws.cell(row=data_header_row, column=col_idx, value=str(column_name))
            
            self.apply_header_style(ws, data_header_row, header_font, header_fill)
            
            # Add DataFrame data (all rows, no limit)
            for row_idx, (_, row_data) in enumerate(data.iterrows(), data_header_row + 1):
                for col_idx, value in enumerate(row_data, 1):
                    # Convert value to string to handle different data types
                    cell_value = str(value) if pd.notna(value) else ""
                    ws.cell(row=row_idx, column=col_idx, value=cell_value)
            
            # Add summary statistics after the data
            summary_start_row = data_header_row + len(data) + 3
            ws.cell(row=summary_start_row, column=1, value="Financial Summary")
            ws.cell(row=summary_start_row, column=1).font = Font(size=14, bold=True)
            
            # Calculate and add summary statistics
            summary_row = summary_start_row + 2
            if 'money' in data.columns:
                total_revenue = data['money'].sum()
                avg_transaction = data['money'].mean()
                max_transaction = data['money'].max()
                min_transaction = data['money'].min()
                transaction_count = len(data)
                
                summary_data = [
                    ["Total Revenue", f"¥{total_revenue:,.2f}"],
                    ["Transaction Count", f"{transaction_count:,}"],
                    ["Average Transaction", f"¥{avg_transaction:.2f}"],
                    ["Highest Transaction", f"¥{max_transaction:.2f}"],
                    ["Lowest Transaction", f"¥{min_transaction:.2f}"]
                ]
                
                for i, (label, value) in enumerate(summary_data, summary_row):
                    ws.cell(row=i, column=1, value=label)
                    ws.cell(row=i, column=2, value=value)
                    ws.cell(row=i, column=1).font = Font(bold=True)
        else:
            # Fallback to sample data if no actual data provided
            headers = ["Category", "Budget", "Actual", "Variance", "Variance %", "Status"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=data_start_row, column=col, value=header)
            
            self.apply_header_style(ws, data_start_row, header_font, header_fill)
            
            sample_data = [
                ["Revenue", 100000, 105000, 5000, "5%", "Above Target"],
                ["Operating Costs", 60000, 58000, -2000, "-3.3%", "Below Budget"],
                ["Marketing", 15000, 16500, 1500, "10%", "Over Budget"],
                ["Net Profit", 25000, 30500, 5500, "22%", "Excellent"]
            ]
            
            for i, row_data in enumerate(sample_data, data_start_row + 1):
                for col, value in enumerate(row_data, 1):
                    ws.cell(row=i, column=col, value=value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column:
                try:
                    # Skip merged cells
                    if hasattr(cell, 'column_letter'):
                        if column_letter is None:
                            column_letter = cell.column_letter
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                except:
                    pass
            if column_letter:
                # Remove width limit to allow full content display
                adjusted_width = max_length + 2
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save file
        filename = f"excel_Finance_Analysis_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.exports_dir, filename)
        wb.save(filepath)
        
        return filepath, filename
    
    def generate_customer_excel(self, data, analysis_text, session_id):
        """Generate Customer-focused Excel file."""
        wb, ws, header_font, header_fill = self.create_workbook("Customer Analysis")
        
        # Title
        ws['A1'] = "Customer Analysis Report"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:F1')
        
        # Date
        ws['A2'] = f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws.merge_cells('A2:F2')
        
        # Analysis Summary
        ws['A4'] = "Customer Analysis Summary"
        ws['A4'].font = Font(size=14, bold=True)
        ws['A5'] = analysis_text if analysis_text else "No analysis available"
        # Dynamically adjust merge range based on content length
        merge_rows = max(8, min(20, len(analysis_text) // 100)) if analysis_text else 8
        ws.merge_cells(f'A5:F{merge_rows}')
        ws['A5'].alignment = Alignment(wrap_text=True, vertical="top")
        
        # Customer Data Structure
        # Adjust starting row based on analysis content length
        row = merge_rows + 3
        headers = ["Customer Segment", "Count", "Revenue", "Avg Order Value", "Retention Rate", "Satisfaction"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header)
        
        self.apply_header_style(ws, row, header_font, header_fill)
        
        # Add sample customer data
        if data is None or (isinstance(data, str) and data == "No data available") or (hasattr(data, 'empty') and data.empty):
            sample_data = [
                ["Premium", 150, 75000, 500, "95%", "4.8/5"],
                ["Standard", 800, 120000, 150, "85%", "4.2/5"],
                ["Basic", 1200, 60000, 50, "70%", "3.9/5"],
                ["Total", 2150, 255000, 119, "83%", "4.3/5"]
            ]
            
            for i, row_data in enumerate(sample_data, row + 1):
                for col, value in enumerate(row_data, 1):
                    ws.cell(row=i, column=col, value=value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column:
                try:
                    # Skip merged cells
                    if hasattr(cell, 'column_letter'):
                        if column_letter is None:
                            column_letter = cell.column_letter
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                except:
                    pass
            if column_letter:
                # Remove width limit to allow full content display
                adjusted_width = max_length + 2
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save file
        filename = f"excel_Customer_Analysis_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.exports_dir, filename)
        wb.save(filepath)
        
        return filepath, filename
    
    def generate_inventory_excel(self, data, analysis_text, session_id):
        """Generate Inventory-focused Excel file."""
        wb, ws, header_font, header_fill = self.create_workbook("Inventory Analysis")
        
        # Title
        ws['A1'] = "Inventory Analysis Report"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:G1')
        
        # Date
        ws['A2'] = f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws.merge_cells('A2:G2')
        
        # Analysis Summary
        ws['A4'] = "Inventory Analysis Summary"
        ws['A4'].font = Font(size=14, bold=True)
        ws['A5'] = analysis_text if analysis_text else "No analysis available"
        # Dynamically adjust merge range based on content length
        merge_rows = max(8, min(20, len(analysis_text) // 100)) if analysis_text else 8
        ws.merge_cells(f'A5:G{merge_rows}')
        ws['A5'].alignment = Alignment(wrap_text=True, vertical="top")
        
        # Inventory Data Structure
        # Adjust starting row based on analysis content length
        row = merge_rows + 3
        headers = ["Product", "Current Stock", "Reorder Level", "Lead Time", "Turnover Rate", "Value", "Status"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header)
        
        self.apply_header_style(ws, row, header_font, header_fill)
        
        # Add sample inventory data
        if data is None or (isinstance(data, str) and data == "No data available") or (hasattr(data, 'empty') and data.empty):
            sample_data = [
                ["Product A", 250, 100, "7 days", "12x/year", "$12,500", "Normal"],
                ["Product B", 80, 150, "14 days", "8x/year", "$8,000", "Low Stock"],
                ["Product C", 500, 200, "5 days", "15x/year", "$25,000", "Overstocked"],
                ["Product D", 120, 100, "10 days", "10x/year", "$6,000", "Normal"]
            ]
            
            for i, row_data in enumerate(sample_data, row + 1):
                for col, value in enumerate(row_data, 1):
                    ws.cell(row=i, column=col, value=value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column:
                try:
                    # Skip merged cells
                    if hasattr(cell, 'column_letter'):
                        if column_letter is None:
                            column_letter = cell.column_letter
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                except:
                    pass
            if column_letter:
                # Remove width limit to allow full content display
                adjusted_width = max_length + 2
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save file
        filename = f"excel_Inventory_Analysis_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.exports_dir, filename)
        wb.save(filepath)
        
        return filepath, filename
    
    def generate_hr_excel(self, data, analysis_text, session_id):
        """Generate HR-focused Excel file."""
        wb, ws, header_font, header_fill = self.create_workbook("HR Analysis")
        
        # Title
        ws['A1'] = "Human Resources Analysis Report"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:F1')
        
        # Date
        ws['A2'] = f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws.merge_cells('A2:F2')
        
        # Analysis Summary
        ws['A4'] = "HR Analysis Summary"
        ws['A4'].font = Font(size=14, bold=True)
        ws['A5'] = analysis_text if analysis_text else "No analysis available"
        # Dynamically adjust merge range based on content length
        merge_rows = max(8, min(20, len(analysis_text) // 100)) if analysis_text else 8
        ws.merge_cells(f'A5:F{merge_rows}')
        ws['A5'].alignment = Alignment(wrap_text=True, vertical="top")
        
        # HR Data Structure
        # Adjust starting row based on analysis content length
        row = merge_rows + 3
        headers = ["Department", "Headcount", "Turnover Rate", "Avg Salary", "Performance", "Satisfaction"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header)
        
        self.apply_header_style(ws, row, header_font, header_fill)
        
        # Add sample HR data
        if data is None or (isinstance(data, str) and data == "No data available") or (hasattr(data, 'empty') and data.empty):
            sample_data = [
                ["Sales", 25, "8%", "$65,000", "4.2/5", "4.1/5"],
                ["Marketing", 12, "12%", "$58,000", "4.0/5", "3.9/5"],
                ["Engineering", 35, "5%", "$85,000", "4.5/5", "4.3/5"],
                ["HR", 8, "10%", "$55,000", "4.1/5", "4.0/5"],
                ["Finance", 15, "6%", "$70,000", "4.3/5", "4.2/5"]
            ]
            
            for i, row_data in enumerate(sample_data, row + 1):
                for col, value in enumerate(row_data, 1):
                    ws.cell(row=i, column=col, value=value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column:
                try:
                    # Skip merged cells
                    if hasattr(cell, 'column_letter'):
                        if column_letter is None:
                            column_letter = cell.column_letter
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                except:
                    pass
            if column_letter:
                # Remove width limit to allow full content display
                adjusted_width = max_length + 2
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save file
        filename = f"excel_HR_Analysis_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.exports_dir, filename)
        wb.save(filepath)
        
        return filepath, filename
    
    def generate_excel_by_type(self, excel_type, data, analysis_text, session_id):
        """Generate Excel file based on the specified type."""
        excel_type = excel_type.lower()
        
        if excel_type == "sales":
            return self.generate_sales_excel(data, analysis_text, session_id)
        elif excel_type == "finance":
            return self.generate_finance_excel(data, analysis_text, session_id)
        elif excel_type == "customer":
            return self.generate_customer_excel(data, analysis_text, session_id)
        elif excel_type == "inventory":
            return self.generate_inventory_excel(data, analysis_text, session_id)
        elif excel_type == "hr":
            return self.generate_hr_excel(data, analysis_text, session_id)
        else:
            # Default to sales if type not recognized
            return self.generate_sales_excel(data, analysis_text, session_id)


def create_excel_file(excel_type, data, analysis_text, session_id):
    """Convenience function to create Excel file."""
    generator = ExcelGenerator()
    return generator.generate_excel_by_type(excel_type, data, analysis_text, session_id)